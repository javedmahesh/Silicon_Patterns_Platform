"""Source readers — turn an engine's output file into a list of plain dict rows.

Supports the formats the four engines actually emit, so an ingestor can take the
richest available source with a uniform interface:

  * ``.csv``           -> Career/Industry CSV exports
  * ``.xlsx``          -> the primary Excel deliverables (openpyxl, imported lazily)
  * ``.jsonl``         -> the Blog Crawler's richest export (one opportunity/line)
  * ``.json``          -> a LinkedIn ``candidates`` export (array of objects, or
                          a Supabase REST dump)

``read_rows`` dispatches on file extension. All readers return ``list[dict]`` with
string keys; values are left as-is (ingestors coerce/normalize). A missing file
raises :class:`FileNotFoundError` — callers decide whether that is fatal.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any


def read_rows(path: str | Path) -> list[dict[str, Any]]:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)
    suffix = p.suffix.lower()
    if suffix == ".csv":
        return read_csv(p)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(p)
    if suffix == ".jsonl":
        return read_jsonl(p)
    if suffix == ".json":
        return read_json(p)
    raise ValueError(f"Unsupported source format: {p.suffix} ({p})")


def read_csv(path: str | Path) -> list[dict[str, Any]]:
    with open(path, "r", newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def read_jsonl(path: str | Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with open(path, "r", encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    return rows


def read_json(path: str | Path) -> list[dict[str, Any]]:
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    # Accept either a bare array or a Supabase-style {"data": [...]} envelope.
    if isinstance(data, dict):
        for key in ("data", "candidates", "rows", "records"):
            if isinstance(data.get(key), list):
                return data[key]
        return [data]
    return list(data)


def read_xlsx(path: str | Path, sheet: str | None = None) -> list[dict[str, Any]]:
    """Read the first (or named) worksheet into dict rows. openpyxl is imported
    here so the rest of P1 (and its tests) run without it installed."""
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:  # pragma: no cover
        raise RuntimeError(
            "Reading .xlsx requires openpyxl (`pip install openpyxl`). "
            "Use a .csv/.jsonl source, or install openpyxl."
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(c) if c is not None else "" for c in next(rows_iter)]
    except StopIteration:
        return []
    out: list[dict[str, Any]] = []
    for raw in rows_iter:
        if raw is None or all(c is None for c in raw):
            continue
        out.append({header[i]: raw[i] for i in range(min(len(header), len(raw)))})
    wb.close()
    return out
